import sys, os
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
sys.path.insert(0, os.path.dirname(__file__))
from docx_design_helpers import insert_in_pPr_order

path = sys.argv[1]
doc = Document(path)

INK = RGBColor(0x1a, 0x1a, 0x1a)
RULE_COLOR = "1A1A1A"
KEEP_FONT_STYLES = {'VerbatimChar', 'Hyperlink'}

# ---- center every figure image + its "Figure N.n: ..." caption paragraph ----
centered = 0
for p in doc.paragraphs:
    xml = p._p.xml
    is_image = '<w:drawing' in xml
    text = p.text.strip()
    is_caption = text.startswith('Figure ') and ':' in text[:40]
    if is_image or is_caption:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        centered += 1


def strip_cell_shading(cell):
    tcPr = cell._tc.find(qn('w:tcPr'))
    if tcPr is not None:
        shd = tcPr.find(qn('w:shd'))
        if shd is not None:
            tcPr.remove(shd)
    for para in cell.paragraphs:
        pPr = para._p.find(qn('w:pPr'))
        if pPr is not None:
            shd = pPr.find(qn('w:shd'))
            if shd is not None:
                pPr.remove(shd)


def set_cell_border(cell, side, sz, color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    el = tcBorders.find(qn(f'w:{side}'))
    if el is None:
        el = OxmlElement(f'w:{side}')
        tcBorders.append(el)
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:color'), color)


def clear_all_cell_borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is not None:
        tcPr.remove(tcBorders)
    tcBorders = OxmlElement('w:tcBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:val'), 'nil')
        tcBorders.append(el)
    tcPr.append(tcBorders)


# ---- redesign real data tables (>1 row) as academic "booktabs" style:
#      no fill, no vertical rules, a rule above the header, a rule below the
#      header, and a rule closing the bottom of the table. The title-page
#      "Submitted By" card is a single-row table and is intentionally left
#      untouched by this loop (it already carries its own gray card styling
#      from the markdown source, and treating it as a "data table" was the
#      bug that double-shaded it blue in the previous build). ----
tables_styled = 0
for table in doc.tables:
    if len(table.rows) < 2:
        continue
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        is_header = r_idx == 0
        for cell in row.cells:
            strip_cell_shading(cell)
            clear_all_cell_borders(cell)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.bold = is_header
                    run.font.color.rgb = INK
            if is_header:
                set_cell_border(cell, 'top', sz=8, color=RULE_COLOR)
                set_cell_border(cell, 'bottom', sz=6, color=RULE_COLOR)
            if r_idx == len(table.rows) - 1:
                set_cell_border(cell, 'bottom', sz=8, color=RULE_COLOR)
    tables_styled += 1

# ---- strip pandoc's auto-generated heading bookmarks — pure clutter, since
#      nothing in this document cross-references them (the TOC is now static
#      text, and Word's native TOC-field, where still used, scans heading
#      styles directly rather than bookmarks) ----
root = doc.element.body
bookmarks_removed = 0
for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
    for el in root.iter(qn(tag)):
        el.getparent().remove(el)
        bookmarks_removed += 1

# ---- font safety net: force every run in the body to Times New Roman
#      unless it deliberately opts into a different font via its style
#      (code spans/blocks use Consolas) ----
font_fixed = 0
for p in doc.paragraphs:
    for run in p.runs:
        rPr = run._element.find(qn('w:rPr'))
        rStyle = rPr.find(qn('w:rStyle')) if rPr is not None else None
        style_id = rStyle.get(qn('w:val')) if rStyle is not None else None
        if style_id in KEEP_FONT_STYLES:
            continue
        if run.font.name and 'Consolas' in run.font.name:
            continue
        if run.font.name != 'Times New Roman':
            run.font.name = 'Times New Roman'
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = OxmlElement('w:rFonts')
                rPr.insert(0, rFonts)
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            rFonts.set(qn('w:cs'), 'Times New Roman')
            font_fixed += 1

doc.save(path)
print(f"{path}: centered {centered} paragraphs, styled {tables_styled} tables, "
      f"removed {bookmarks_removed} bookmarks, fixed {font_fixed} stray fonts")
