import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from docx_design_helpers import (
    style_bottom_border, style_left_accent_bar, set_pPr_border,
    rStyle_shade, style_pPr_shade, style_box_border, style_indent,
    force_doc_default_font,
)

# Monochrome document chrome — only images/diagrams carry color. Rules and
# bars use the same ink tone as body text, not an accent color.
INK = RGBColor(0x1a, 0x1a, 0x1a)
INK_HEX = "1A1A1A"
GRAY_LINE = "999999"
GRAY_FILL = "F2F2F2"

SRC = "markdown/build/base-reference.docx"
OUT = "markdown/build/nexis-reference-compact.docx"

doc = Document(SRC)
FONT = "Times New Roman"

def set_font(style, name=FONT, size=None, bold=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

force_doc_default_font(doc, name=FONT)

normal = doc.styles['Normal']
set_font(normal, size=10.5, color=INK)
pf = normal.paragraph_format
pf.line_spacing = 1.15
pf.space_after = Pt(5)
pf.space_before = Pt(0)

h1 = doc.styles['Heading1']
set_font(h1, size=15, bold=True, color=INK)
h1.paragraph_format.space_before = Pt(0)
h1.paragraph_format.space_after = Pt(10)
h1.paragraph_format.page_break_before = True
h1.paragraph_format.keep_with_next = True
style_bottom_border(h1, sz=12, color=INK_HEX, space=5)

h2 = doc.styles['Heading2']
set_font(h2, size=12.5, bold=True, color=INK)
h2.paragraph_format.space_before = Pt(10)
h2.paragraph_format.space_after = Pt(4)
h2.paragraph_format.page_break_before = False
h2.paragraph_format.keep_with_next = True

h3 = doc.styles['Heading3']
set_font(h3, size=11, bold=True, color=INK)
h3.paragraph_format.space_before = Pt(7)
h3.paragraph_format.space_after = Pt(3)
h3.paragraph_format.keep_with_next = True

h4 = doc.styles['Heading4']
set_font(h4, size=10.5, bold=True, color=INK)
h4.paragraph_format.space_before = Pt(5)
h4.paragraph_format.space_after = Pt(2)

title = doc.styles['Title']
set_font(title, size=22, bold=True, color=INK)
title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

toch = doc.styles['TOCHeading']
set_font(toch, size=16, bold=True, color=INK)
toch.paragraph_format.page_break_before = True
toch.paragraph_format.space_after = Pt(12)

cap = doc.styles['Caption']
set_font(cap, size=9, color=INK)
cap.paragraph_format.space_before = Pt(2)
cap.paragraph_format.space_after = Pt(8)

vchar = doc.styles['Verbatim Char']
set_font(vchar, name='Consolas', size=8.5, color=INK)
rStyle_shade(vchar, fill=GRAY_FILL)

from docx.enum.style import WD_STYLE_TYPE
try:
    sc = doc.styles['Source Code']
except KeyError:
    sc = doc.styles.add_style('Source Code', WD_STYLE_TYPE.PARAGRAPH)
    sc.base_style = doc.styles['Normal']
sc.paragraph_format.line_spacing = 1.0
sc.paragraph_format.space_before = Pt(6)
sc.paragraph_format.space_after = Pt(9)
style_pPr_shade(sc, fill=GRAY_FILL)
style_box_border(sc, sz=4, color=GRAY_LINE, space=6)
style_indent(sc, left_twips=140, right_twips=140)

if 'Table' in [s.name for s in doc.styles]:
    tbl = doc.styles['Table']
    set_font(tbl, size=9.5, color=INK)

if 'Compact' in [s.name for s in doc.styles]:
    cm = doc.styles['Compact']
    set_font(cm, size=9.5, color=INK)
    cm.paragraph_format.space_after = Pt(2)

# Hyperlinks: black + underline (not blue) — this is a monochrome document,
# links are distinguished by underline the way print-friendly docs do it.
if 'Hyperlink' in [s.name for s in doc.styles]:
    hl = doc.styles['Hyperlink']
    hl.font.underline = True
    hl.font.color.rgb = INK

for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.different_first_page_header_footer = True

# No decorative running header — plain academic report, header left blank.
# Footer: plain centered page number, suppressed on the title page.
section = doc.sections[0]

for footer_obj, is_first in ((section.first_page_footer, True), (section.footer, False)):
    footer_obj.is_linked_to_previous = False
    if is_first:
        # leave the title page's footer empty
        continue
    p = footer_obj.paragraphs[0] if footer_obj.paragraphs else footer_obj.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run_el in list(p.runs):
        run_el._element.getparent().remove(run_el._element)

    def add_field(paragraph, instr_text):
        run = paragraph.add_run()
        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = instr_text
        fld_sep = OxmlElement('w:fldChar')
        fld_sep.set(qn('w:fldCharType'), 'separate')
        fld_end = OxmlElement('w:fldChar')
        fld_end.set(qn('w:fldCharType'), 'end')
        r_el = run._element
        r_el.append(fld_begin)
        r2 = paragraph.add_run()._element
        r2.append(instr)
        r3 = paragraph.add_run()._element
        r3.append(fld_sep)
        r4 = paragraph.add_run()._element
        r4.append(fld_end)

    run = p.add_run()
    run.font.name = FONT
    run.font.size = Pt(9.5)
    run.font.color.rgb = INK
    add_field(p, "PAGE")

doc.save(OUT)
print("wrote", OUT)
